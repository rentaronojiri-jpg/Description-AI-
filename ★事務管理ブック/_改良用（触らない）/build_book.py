import sys, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.chart import BarChart, Reference

FONT = "游ゴシック"
NAVY = "1F3864"
IN_FILL = PatternFill("solid", fgColor="FFF9E5")   # 入力セル（黄色）
AUTO_FILL = PatternFill("solid", fgColor="F2F2F2")  # 自動計算（グレー）
HDR_FILL = PatternFill("solid", fgColor=NAVY)
SUB_FILL = PatternFill("solid", fgColor="D9E1F2")
thin = Side(style="thin", color="A6A6A6")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
YEN = '#,##0"円"'
NUM = '#,##0'
DATE = 'yyyy/m/d'

# 行範囲
T0, T1 = 6, 105     # 案件台帳
B0, B1 = 6, 205     # 予算管理
P0, P1 = 11, 210     # 支払管理
C0, C1 = 6, 505     # 入出金帳

S_SET, S_LED, S_EST, S_BUD, S_INV, S_PAY, S_CASH, S_DASH, S_HELP = (
    "設定", "案件台帳", "見積書", "予算管理", "請求書", "支払管理", "入出金帳", "ダッシュボード", "使い方")


def q(s):
    return "'" + s + "'"


TAX = f"{q(S_SET)}!$C$11"
SITE = f"{q(S_SET)}!$C$12"
VALID = f"{q(S_SET)}!$C$13"
START = f"{q(S_SET)}!$C$14"
OPEN_BAL = f"{q(S_SET)}!$C$15"
LED_NO = f"{q(S_LED)}!$A${T0}:$A${T1}"


def led_col(c):
    return f"{q(S_LED)}!${c}${T0}:${c}${T1}"


def lookup(key, col):
    return f'IF({key}="","",IFERROR(INDEX({led_col(col)},MATCH({key},{LED_NO},0)),""))'


def st(c, fill=None, fmt=None, bold=False, color=None, size=11, h=None, v="center", wrap=False, border=True):
    c.font = Font(name=FONT, size=size, bold=bold, color=color)
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(horizontal=h, vertical=v, wrap_text=wrap)
    if border:
        c.border = BORDER


def put(ws, ref, val=None, kind=None, fmt=None, **kw):
    c = ws[ref]
    if val is not None:
        c.value = val
    fill = kw.pop("fill", None) or (IN_FILL if kind == "in" else AUTO_FILL if kind == "auto" else None)
    st(c, fill=fill, fmt=fmt, **kw)
    return c


def header(ws, row, labels, col0=1):
    for i, t in enumerate(labels):
        c = ws.cell(row=row, column=col0 + i, value=t)
        st(c, fill=HDR_FILL, bold=True, color="FFFFFF", h="center", wrap=True)
    ws.row_dimensions[row].height = 30


def widths(ws, ws_widths):
    for k, w in ws_widths.items():
        ws.column_dimensions[k].width = w


def title(ws, text, sub=None, span="A1:H1"):
    ws[span.split(":")[0]].value = text
    ws[span.split(":")[0]].font = Font(name=FONT, size=16, bold=True, color=NAVY)
    if sub:
        ws["A2"].value = sub
        ws["A2"].font = Font(name=FONT, size=10, color="595959")


def dv_list(ws, rng, src, msg=None):
    dv = DataValidation(type="list", formula1=src, allow_blank=True)
    dv.error = "リストから選んでください"
    dv.showErrorMessage = True
    ws.add_data_validation(dv)
    dv.add(rng)


def page(ws, landscape=True, area=None):
    ws.page_setup.paperSize = 8  # A3
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1 if area else 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    if area:
        ws.print_area = area


def build(sample, path):
    wb = Workbook()
    ws0 = wb.active
    ws0.title = S_HELP
    ws_set = wb.create_sheet(S_SET)
    ws_led = wb.create_sheet(S_LED)
    ws_est = wb.create_sheet(S_EST)
    ws_bud = wb.create_sheet(S_BUD)
    ws_inv = wb.create_sheet(S_INV)
    ws_pay = wb.create_sheet(S_PAY)
    ws_cash = wb.create_sheet(S_CASH)
    ws_dash = wb.create_sheet(S_DASH)
    wb.move_sheet(S_DASH, offset=-7)   # ダッシュボードを2番目へ
    for w, col in [(ws0, "808080"), (ws_set, "808080"), (ws_led, NAVY), (ws_est, "2E75B6"), (ws_bud, "548235"),
                   (ws_inv, "2E75B6"), (ws_pay, "C55A11"), (ws_cash, "C55A11"), (ws_dash, "BF9000")]:
        w.sheet_properties.tabColor = col
        w.sheet_view.showGridLines = False

    # ===================== 設定 =====================
    ws = ws_set
    title(ws, "設定（自社情報・リスト）", "黄色のセルを書き換えてください。他のシートに自動で反映されます。", "A1:E1")
    widths(ws, {"A": 3, "B": 24, "C": 40, "D": 24, "E": 20})
    info = [
        (4, "会社名", "株式会社サンプル工務店" if sample else None, None),
        (5, "代表者名", "代表取締役 山田 太郎" if sample else None, None),
        (6, "郵便番号", "000-0000" if sample else None, None),
        (7, "住所", "○○県○○市サンプル町1-2-3" if sample else None, None),
        (8, "電話番号", "000-0000-0000" if sample else None, None),
        (9, "インボイス登録番号", "T0000000000000" if sample else None, None),
        (10, "振込先", "サンプル銀行 ○○支店 普通 0000000 株式会社サンプル工務店" if sample else None, None),
        (11, "消費税率", 0.10, "0%"),
        (12, "支払サイト（請求日から何日後）", 30, '0"日"'),
        (13, "見積の有効期限（何日）", 30, '0"日"'),
        (14, "集計の開始月（ダッシュボード）", dt.date(2026, 7, 1) if sample else dt.date(dt.date.today().year, 1, 1), 'yyyy"年"m"月"'),
        (15, "入出金帳の繰越残高", 1000000 if sample else 0, YEN),
    ]
    for r, lab, val, fmt in info:
        put(ws, f"B{r}", lab, fill=SUB_FILL, bold=True)
        put(ws, f"C{r}", val, "in", fmt)
    ws["B3"].value = "■ 自社情報・基本設定"
    ws["B3"].font = Font(name=FONT, bold=True, color=NAVY)
    ws["B17"].value = "■ ドロップダウンの選択肢（増やす・書き換えOK。空き行に追記できます）"
    ws["B17"].font = Font(name=FONT, bold=True, color=NAVY)
    header(ws, 18, ["案件の状況", "工種", "入出金の科目", "支払状況（変更しない）"], col0=2)
    status = ["見積中", "受注", "施工中", "完工", "入金済み", "失注"]
    kinds = ["解体", "仮設・足場", "土工・基礎", "鉄筋・型枠", "大工・木工", "内装", "塗装", "防水", "左官", "電気",
             "給排水・設備", "屋根・板金", "外構", "清掃・廃材処分", "諸経費", "その他"]
    accts = ["請負代金入金", "その他入金", "外注費", "材料費", "交通費（ETC・ガソリン）", "駐車場", "消耗品", "事務所経費",
             "給与・賞与", "税金・保険", "その他出金"]
    pays = ["未払い", "問題あり", "支払い済み"]
    for i in range(20):
        r = 19 + i
        for col, lst in zip("BCDE", [status, kinds, accts, pays]):
            put(ws, f"{col}{r}", lst[i] if i < len(lst) else None, "auto" if col == "E" else "in")
    LST_STATUS = f"{q(S_SET)}!$B$19:$B$38"
    LST_KIND = f"{q(S_SET)}!$C$19:$C$38"
    LST_ACCT = f"{q(S_SET)}!$D$19:$D$38"
    LST_PAY = f"{q(S_SET)}!$E$19:$E$21"
    page(ws, landscape=False)

    # ===================== 案件台帳 =====================
    ws = ws_led
    title(ws, "案件台帳", "案件を1行ずつ登録します。ここに登録した「案件No」を、他のシートで選びます。　黄＝入力　灰＝自動計算", "A1:N1")
    widths(ws, {"A": 11, "B": 34, "C": 24, "D": 26, "E": 10, "F": 15, "G": 13, "H": 15, "I": 15, "J": 15, "K": 9,
                "L": 15, "M": 15, "N": 24})
    header(ws, 5, ["案件No", "案件名", "取引先", "現場住所", "状況", "請負金額\n（税抜）", "消費税", "請負金額\n（税込）",
                   "外注費計\n（税抜）", "粗利\n（税抜）", "粗利率", "入金済額\n（税込）", "未入金額\n（税込）", "備考"])
    ws.merge_cells("A3:E3")
    put(ws, "A3", "合計（見積中・失注は除く）", fill=SUB_FILL, bold=True, h="right")
    for col in "FGHI":
        put(ws, f"{col}3", f'=SUM({col}{T0}:{col}{T1})-SUMIF($E${T0}:$E${T1},"見積中",{col}{T0}:{col}{T1})-SUMIF($E${T0}:$E${T1},"失注",{col}{T0}:{col}{T1})', "auto", YEN, bold=True)
    for col in "JLM":
        put(ws, f"{col}3", f"=SUM({col}{T0}:{col}{T1})", "auto", YEN, bold=True)
    put(ws, "K3", '=IF(N(F3)=0,"",J3/F3)', "auto", "0.0%", bold=True)
    for r in range(T0, T1 + 1):
        for col in "ABCDEN":
            put(ws, f"{col}{r}", None, "in", None)
        put(ws, f"F{r}", None, "in", YEN)
        put(ws, f"G{r}", f'=IF(F{r}="","",ROUND(F{r}*{TAX},0))', "auto", YEN)
        put(ws, f"H{r}", f'=IF(F{r}="","",F{r}+G{r})', "auto", YEN)
        put(ws, f"I{r}", f'=IF(A{r}="","",SUMPRODUCT(({q(S_BUD)}!$A${B0}:$A${B1}=A{r})*1,{q(S_BUD)}!$E${B0}:$E${B1}))', "auto", YEN)
        put(ws, f"J{r}", f'=IF(OR(A{r}="",F{r}="",E{r}="見積中",E{r}="失注"),"",F{r}-I{r})', "auto", YEN)
        put(ws, f"K{r}", f'=IF(OR(J{r}="",N(F{r})=0),"",J{r}/F{r})', "auto", "0.0%")
        put(ws, f"L{r}", f'=IF(A{r}="","",SUMPRODUCT(({q(S_CASH)}!$B${C0}:$B${C1}=A{r})*1,{q(S_CASH)}!$G${C0}:$G${C1}))', "auto", YEN)
        put(ws, f"M{r}", f'=IF(OR(A{r}="",H{r}="",E{r}="見積中",E{r}="失注"),"",H{r}-L{r})', "auto", YEN)
    if sample:
        led = [("2610-01", "サンプルマンション 外壁改修工事", "サンプル不動産株式会社", "○○県○○市サンプル町1-1", "施工中", 8000000),
               ("2610-02", "サンプル店舗 内装工事", "サンプル商事株式会社", "○○県○○市サンプル町2-2", "受注", 3500000),
               ("2609-03", "サンプル邸 屋根防水工事", "個人（サンプル様）", "○○県○○市サンプル町3-3", "完工", 1200000),
               ("2610-04", "サンプルビル 解体工事", "サンプル開発株式会社", "○○県○○市サンプル町4-4", "見積中", 2400000)]
        for i, row in enumerate(led):
            for col, val in zip("ABCDEF", row):
                ws[f"{col}{T0 + i}"].value = val
    dv_list(ws, f"E{T0}:E{T1}", LST_STATUS)
    ws.freeze_panes = "C6"
    ws.auto_filter.ref = f"A5:N{T1}"
    ws.conditional_formatting.add(f"E{T0}:E{T1}", FormulaRule(formula=[f'$E{T0}="入金済み"'], fill=PatternFill("solid", bgColor="C6E0B4")))
    ws.conditional_formatting.add(f"E{T0}:E{T1}", FormulaRule(formula=[f'$E{T0}="施工中"'], fill=PatternFill("solid", bgColor="BDD7EE")))
    ws.conditional_formatting.add(f"E{T0}:E{T1}", FormulaRule(formula=[f'$E{T0}="失注"'], fill=PatternFill("solid", bgColor="D9D9D9")))
    ws.conditional_formatting.add(f"M{T0}:M{T1}", FormulaRule(formula=[f'AND(ISNUMBER($M{T0}),$M{T0}>0,OR($E{T0}="完工",$E{T0}="入金済み"))'], font=Font(color="C00000", bold=True)))
    page(ws)

    # ===================== 見積書 =====================
    ws = ws_est
    widths(ws, {"A": 6, "B": 17, "C": 38, "D": 9, "E": 7, "F": 13, "G": 16, "H": 18})
    ws.merge_cells("A1:H1")
    put(ws, "A1", "御　見　積　書", size=22, bold=True, h="center", border=False)
    ws.row_dimensions[1].height = 40
    put(ws, "A3", "案件No", fill=SUB_FILL, bold=True, h="center")
    put(ws, "B3", "2610-02" if sample else None, "in", h="center")
    put(ws, "F3", "見積番号", fill=SUB_FILL, bold=True, h="center")
    ws.merge_cells("G3:H3")
    put(ws, "G3", '=IF(B3="","",B3&"-E")', "auto", h="center")
    put(ws, "H3", None, "auto")
    put(ws, "A4", "宛名", fill=SUB_FILL, bold=True, h="center")
    ws.merge_cells("B4:D4")
    put(ws, "B4", f'=IF(B3="","",{lookup("B3", "C")}&" 御中")', "auto", bold=True, size=13)
    for c in "CD":
        put(ws, f"{c}4", None, "auto")
    put(ws, "F4", "見積日", fill=SUB_FILL, bold=True, h="center")
    ws.merge_cells("G4:H4")
    put(ws, "G4", dt.date(2026, 10, 1) if sample else None, "in", DATE, h="center")
    put(ws, "H4", None, "in")
    put(ws, "A5", "件名", fill=SUB_FILL, bold=True, h="center")
    ws.merge_cells("B5:D5")
    put(ws, "B5", "=" + lookup("B3", "B"), "auto")
    for c in "CD":
        put(ws, f"{c}5", None, "auto")
    put(ws, "F5", "有効期限", fill=SUB_FILL, bold=True, h="center")
    ws.merge_cells("G5:H5")
    put(ws, "G5", f'=IF(G4="","",G4+{VALID})', "auto", DATE, h="center")
    put(ws, "H5", None, "auto")
    put(ws, "A6", "現場", fill=SUB_FILL, bold=True, h="center")
    ws.merge_cells("B6:D6")
    put(ws, "B6", "=" + lookup("B3", "D"), "auto")
    for c in "CD":
        put(ws, f"{c}6", None, "auto")
    # 金額と自社欄
    ws.merge_cells("A8:B8")
    put(ws, "A8", "お見積金額（税込）", fill=SUB_FILL, bold=True, h="center")
    put(ws, "B8", None, fill=SUB_FILL)
    ws.merge_cells("C8:D8")
    put(ws, "C8", '=IF(G45=0,"",G47)', "auto", YEN, bold=True, size=14, h="center")
    put(ws, "D8", None, "auto")
    comp = [(8, f"={q(S_SET)}!C4", True), (9, f'=IF({q(S_SET)}!C6="","","〒"&{q(S_SET)}!C6&" ")&{q(S_SET)}!C7', False),
            (10, f'=IF({q(S_SET)}!C8="","","TEL "&{q(S_SET)}!C8)', False), (11, f'=IF({q(S_SET)}!C9="","","登録番号 "&{q(S_SET)}!C9)', False), (12, f"={q(S_SET)}!C5", False)]
    for r, f, b in comp:
        ws.merge_cells(f"F{r}:H{r}")
        put(ws, f"F{r}", f, bold=b, border=False, size=11 if not b else 12)
    header(ws, 14, ["No", "工種", "品名・仕様", "数量", "単位", "単価", "金額", "備考"])
    D0, D1 = 15, 44
    est_rows = [
        ("解体", "既存内装解体・処分", 1, "式", 350000),
        ("内装", "軽量鉄骨下地・ボード張り", 85, "㎡", 6500),
        ("内装", "クロス張り", 210, "㎡", 1800),
        ("内装", "床 長尺シート張り", 60, "㎡", 5500),
        ("電気", "照明・コンセント工事", 1, "式", 780000),
        ("諸経費", "現場管理費・諸経費", 1, "式", 1109500),
    ]
    for i, r in enumerate(range(D0, D1 + 1)):
        put(ws, f"A{r}", f'=IF(C{r}="","",{i + 1})', "auto", h="center")
        for col in "BCEH":
            put(ws, f"{col}{r}", None, "in")
        put(ws, f"D{r}", None, "in", 'General')
        put(ws, f"F{r}", None, "in", NUM)
        put(ws, f"G{r}", f'=IF(OR(D{r}="",F{r}=""),"",ROUND(D{r}*F{r},0))', "auto", NUM)
        if sample and i < len(est_rows):
            k, n, qty, u, p = est_rows[i]
            ws[f"B{r}"].value, ws[f"C{r}"].value, ws[f"D{r}"].value, ws[f"E{r}"].value, ws[f"F{r}"].value = k, n, qty, u, p
    dv_list(ws, f"B{D0}:B{D1}", LST_KIND)
    dv_list(ws, "B3", LED_NO)
    ws.merge_cells("A45:F45")
    put(ws, "A45", "小計（税抜）", fill=SUB_FILL, bold=True, h="right")
    put(ws, "G45", f"=SUM(G{D0}:G{D1})", "auto", YEN, bold=True)
    ws.merge_cells("A46:F46")
    put(ws, "A46", "消費税", fill=SUB_FILL, bold=True, h="right")
    put(ws, "G46", f"=ROUND(G45*{TAX},0)", "auto", YEN, bold=True)
    ws.merge_cells("A47:F47")
    put(ws, "A47", "合計（税込）", fill=SUB_FILL, bold=True, h="right")
    put(ws, "G47", "=G45+G46", "auto", YEN, bold=True, size=13)
    put(ws, "A49", "備考", fill=SUB_FILL, bold=True, h="center")
    ws.merge_cells("B49:H52")
    put(ws, "B49", "お支払条件・工期・有効期限などをここに記入" if sample else None, "in", v="top", wrap=True)
    page(ws, landscape=False, area="A1:H52")

    # ===================== 予算管理 =====================
    ws = ws_bud
    title(ws, "予算管理（外注・発注の一覧）", "業者への発注を1件1行で入力します。案件ごとの外注費・粗利は「案件台帳」に自動集計されます。見積のある項目だけ入力してください。", "A1:H1")
    widths(ws, {"A": 11, "B": 32, "C": 16, "D": 26, "E": 15, "F": 15, "G": 13, "H": 28})
    header(ws, 5, ["案件No", "案件名", "工種", "業者名", "発注金額\n（税抜）", "発注金額\n（税込）", "発注日", "備考"])
    put(ws, "D3", "合計", fill=SUB_FILL, bold=True, h="center")
    put(ws, "E3", f"=SUM(E{B0}:E{B1})", "auto", YEN, bold=True)
    put(ws, "F3", f"=SUM(F{B0}:F{B1})", "auto", YEN, bold=True)
    bud = [("2610-01", "仮設・足場", "サンプル足場工業", 1200000, dt.date(2026, 9, 20)),
           ("2610-01", "塗装", "サンプル塗装店", 2800000, dt.date(2026, 9, 25)),
           ("2610-01", "防水", "サンプル防水工業", 900000, dt.date(2026, 10, 1)),
           ("2610-02", "内装", "サンプル内装", 1300000, dt.date(2026, 10, 5)),
           ("2610-02", "電気", "サンプル電設", 400000, dt.date(2026, 10, 8)),
           ("2609-03", "防水", "サンプル防水工業", 500000, dt.date(2026, 9, 1)),
           ("2609-03", "屋根・板金", "サンプル板金", 300000, dt.date(2026, 9, 3))]
    for i, r in enumerate(range(B0, B1 + 1)):
        put(ws, f"A{r}", None, "in")
        put(ws, f"B{r}", "=" + lookup(f"A{r}", "B"), "auto")
        for col in "CDH":
            put(ws, f"{col}{r}", None, "in")
        put(ws, f"E{r}", None, "in", YEN)
        put(ws, f"F{r}", f'=IF(E{r}="","",ROUND(E{r}*(1+{TAX}),0))', "auto", YEN)
        put(ws, f"G{r}", None, "in", DATE)
        if sample and i < len(bud):
            a, k, v, amt, d = bud[i]
            ws[f"A{r}"].value, ws[f"C{r}"].value, ws[f"D{r}"].value, ws[f"E{r}"].value, ws[f"G{r}"].value = a, k, v, amt, d
    dv_list(ws, f"A{B0}:A{B1}", LED_NO)
    dv_list(ws, f"C{B0}:C{B1}", LST_KIND)
    ws.freeze_panes = "C6"
    ws.auto_filter.ref = f"A5:H{B1}"
    page(ws)

    # ===================== 請求書 =====================
    ws = ws_inv
    widths(ws, {"A": 16, "B": 40, "C": 10, "D": 10, "E": 18, "F": 22})
    ws.merge_cells("A1:F1")
    put(ws, "A1", "請　求　書", size=22, bold=True, h="center", border=False)
    ws.row_dimensions[1].height = 40
    put(ws, "A3", "案件No", fill=SUB_FILL, bold=True, h="center")
    put(ws, "B3", "2609-03" if sample else None, "in", h="center")
    put(ws, "E3", "請求書番号", fill=SUB_FILL, bold=True, h="center")
    put(ws, "F3", f'=IF(OR(B3="",F4=""),"",B3&"-"&TEXT(F4,"yymmdd"))', "auto", h="center")
    put(ws, "A4", "宛名", fill=SUB_FILL, bold=True, h="center")
    put(ws, "B4", f'=IF(B3="","",{lookup("B3", "C")}&" 御中")', "auto", bold=True, size=13)
    put(ws, "E4", "請求日", fill=SUB_FILL, bold=True, h="center")
    put(ws, "F4", dt.date(2026, 9, 30) if sample else None, "in", DATE, h="center")
    put(ws, "A5", "件名", fill=SUB_FILL, bold=True, h="center")
    put(ws, "B5", "=" + lookup("B3", "B"), "auto")
    put(ws, "E5", "お支払期限", fill=SUB_FILL, bold=True, h="center")
    put(ws, "F5", f'=IF(F4="","",F4+{SITE})', "auto", DATE, h="center")
    put(ws, "A6", "現場", fill=SUB_FILL, bold=True, h="center")
    put(ws, "B6", "=" + lookup("B3", "D"), "auto")
    put(ws, "E6", "請求割合", fill=SUB_FILL, bold=True, h="center")
    put(ws, "F6", 1, "in", "0%", h="center")
    ws.merge_cells("A8:B8")
    put(ws, "A8", "ご請求金額（税込）", fill=SUB_FILL, bold=True, h="center")
    put(ws, "B8", None, fill=SUB_FILL)
    ws.merge_cells("C8:D8")
    put(ws, "C8", '=IF(F19="","",F19)', "auto", YEN, bold=True, size=14, h="center")
    put(ws, "D8", None, "auto")
    for r, f, b in [(8, f"={q(S_SET)}!C4", True), (9, f'=IF({q(S_SET)}!C6="","","〒"&{q(S_SET)}!C6&" ")&{q(S_SET)}!C7', False),
                    (10, f'=IF({q(S_SET)}!C8="","","TEL "&{q(S_SET)}!C8)', False), (11, f'=IF({q(S_SET)}!C9="","","登録番号 "&{q(S_SET)}!C9)', False), (12, f"={q(S_SET)}!C5", False)]:
        ws.merge_cells(f"E{r}:F{r}")
        put(ws, f"E{r}", f, bold=(r == 8), border=False)
    header(ws, 14, ["品名", "", "数量", "単位", "単価", "金額"])
    ws.merge_cells("A14:B14")
    ws.merge_cells("A15:B15")
    put(ws, "A15", f'=IF(B5="","",B5&" 一式")', "in")
    put(ws, "B15", None, "in")
    put(ws, "C15", f'=IF(B3="","",1)', "auto", h="center")
    put(ws, "D15", f'=IF(B3="","","式")', "auto", h="center")
    put(ws, "E15", f'=IF(B3="","",ROUND(IFERROR(INDEX({led_col("F")},MATCH(B3,{LED_NO},0)),0)*F6,0))', "auto", NUM)
    put(ws, "F15", '=IF(E15="","",E15)', "auto", NUM)
    ws.merge_cells("A17:E17")
    put(ws, "A17", "小計（税抜）", fill=SUB_FILL, bold=True, h="right")
    put(ws, "F17", '=N(F15)', "auto", YEN, bold=True)
    ws.merge_cells("A18:E18")
    put(ws, "A18", "消費税", fill=SUB_FILL, bold=True, h="right")
    put(ws, "F18", f"=ROUND(F17*{TAX},0)", "auto", YEN, bold=True)
    ws.merge_cells("A19:E19")
    put(ws, "A19", "合計（税込）", fill=SUB_FILL, bold=True, h="right")
    put(ws, "F19", '=IF(B3="","",F17+F18)', "auto", YEN, bold=True, size=13)
    put(ws, "A21", "お振込先", fill=SUB_FILL, bold=True, h="center")
    ws.merge_cells("B21:F21")
    put(ws, "B21", f"={q(S_SET)}!C10", "auto")
    put(ws, "A23", "備考", fill=SUB_FILL, bold=True, h="center")
    ws.merge_cells("B23:F26")
    put(ws, "B23", "お振込手数料はご負担をお願いいたします。" if sample else None, "in", v="top", wrap=True)
    dv_list(ws, "B3", LED_NO)
    page(ws, landscape=False, area="A1:F26")

    # ===================== 支払管理 =====================
    ws = ws_pay
    title(ws, "支払管理（業者への支払い）", "業者から届いた請求書を1枚1行で入力。支払状況を選ぶと行の色が変わります。期限を過ぎた未払いは赤字になります。", "A1:L1")
    widths(ws, {"A": 6, "B": 24, "C": 11, "D": 30, "E": 14, "F": 13, "G": 15, "H": 13, "I": 12, "J": 13, "K": 24, "L": 11})
    header(ws, 10, ["No", "支払先（業者名）", "案件No", "案件名", "請求書番号", "請求日", "請求金額\n（税込）", "支払期限",
                   "支払状況", "支払日", "備考", "期限まで\n（日）"])
    header(ws, 4, ["支払状況", "件数", "金額"], col0=2)
    for i, s in enumerate(pays):
        r = 5 + i
        put(ws, f"B{r}", s, fill=SUB_FILL, bold=True)
        put(ws, f"C{r}", f'=COUNTIF($I${P0}:$I${P1},B{r})', "auto", '0"件"')
        put(ws, f"D{r}", f'=SUMIF($I${P0}:$I${P1},B{r},$G${P0}:$G${P1})', "auto", YEN)
    put(ws, "B8", "合計", fill=SUB_FILL, bold=True)
    put(ws, "C8", "=SUM(C5:C7)", "auto", '0"件"', bold=True)
    put(ws, "D8", "=SUM(D5:D7)", "auto", YEN, bold=True)
    pay = [("サンプル足場工業", "2610-01", "A-101", dt.date(2026, 9, 30), 1320000, dt.date(2026, 10, 31), "未払い", None),
           ("サンプル塗装店", "2610-01", "S-55", dt.date(2026, 9, 30), 3080000, dt.date(2026, 10, 31), "未払い", None),
           ("サンプル防水工業", "2609-03", "B-20", dt.date(2026, 9, 30), 550000, dt.date(2026, 10, 10), "未払い", None),
           ("サンプル板金", "2609-03", "K-8", dt.date(2026, 9, 30), 330000, dt.date(2026, 10, 10), "支払い済み", dt.date(2026, 10, 9)),
           ("サンプル内装", "2610-02", "N-3", dt.date(2026, 10, 10), 1430000, dt.date(2026, 11, 30), "問題あり", None)]
    for i, r in enumerate(range(P0, P1 + 1)):
        put(ws, f"A{r}", f'=IF(B{r}="","",ROW()-{P0 - 1})', "auto", h="center")
        for col in "BCEK":
            put(ws, f"{col}{r}", None, "in")
        put(ws, f"D{r}", "=" + lookup(f"C{r}", "B"), "auto")
        put(ws, f"F{r}", None, "in", DATE)
        put(ws, f"G{r}", None, "in", YEN)
        put(ws, f"H{r}", None, "in", DATE)
        put(ws, f"I{r}", None, "in", h="center")
        put(ws, f"J{r}", None, "in", DATE)
        put(ws, f"L{r}", f'=IF(OR(H{r}="",I{r}="支払い済み",B{r}=""),"",H{r}-TODAY())', "auto", '0;[Red]-0')
        if sample and i < len(pay):
            b, c, e, f, g, h, s, j = pay[i]
            for col, val in zip("BCEFGHIJ", [b, c, e, f, g, h, s, j]):
                ws[f"{col}{r}"].value = val
    dv_list(ws, f"C{P0}:C{P1}", LED_NO)
    dv_list(ws, f"I{P0}:I{P1}", LST_PAY)
    rng = f"A{P0}:L{P1}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'$I{P0}="未払い"'], fill=PatternFill("solid", bgColor="FFF2CC")))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'$I{P0}="問題あり"'], fill=PatternFill("solid", bgColor="F8CBAD")))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'$I{P0}="支払い済み"'], fill=PatternFill("solid", bgColor="C6E0B4")))
    ws.conditional_formatting.add(f"H{P0}:H{P1}", FormulaRule(formula=[f'AND($H{P0}<>"",$B{P0}<>"",$I{P0}<>"支払い済み",$H{P0}<TODAY())'], font=Font(color="C00000", bold=True)))
    ws.freeze_panes = f"C{P0}"
    ws.auto_filter.ref = f"A10:L{P1}"
    page(ws)

    # ===================== 入出金帳 =====================
    ws = ws_cash
    title(ws, "入出金帳", "お金が動いたら1行ずつ。入金か出金のどちらか一方に金額を入れます（税込）。案件に関係するものは案件Noを選びます。", "A1:I1")
    widths(ws, {"A": 13, "B": 11, "C": 30, "D": 24, "E": 24, "F": 28, "G": 15, "H": 15, "I": 16})
    put(ws, "A2", "繰越残高", fill=SUB_FILL, bold=True, h="center")
    put(ws, "B2", f"={OPEN_BAL}", "auto", YEN)
    ws.merge_cells("B2:C2")
    put(ws, "C2", None, "auto")
    put(ws, "A3", "合計", fill=SUB_FILL, bold=True, h="center")
    put(ws, "G3", f"=SUM(G{C0}:G{C1})", "auto", YEN, bold=True)
    put(ws, "H3", f"=SUM(H{C0}:H{C1})", "auto", YEN, bold=True)
    put(ws, "F3", "現在の残高 →", fill=SUB_FILL, bold=True, h="right")
    put(ws, "I3", f"={OPEN_BAL}+G3-H3", "auto", YEN, bold=True)
    header(ws, 5, ["日付", "案件No", "案件名", "科目", "相手先", "摘要", "入金（税込）", "出金（税込）", "残高"])
    cash = [(dt.date(2026, 9, 15), None, "事務所経費", "サンプル不動産管理", "事務所家賃", None, 55000),
            (dt.date(2026, 9, 30), "2609-03", "請負代金入金", "個人（サンプル様）", "屋根防水工事 代金", 1320000, None),
            (dt.date(2026, 10, 1), "2610-01", "請負代金入金", "サンプル不動産株式会社", "外壁改修 着手金", 2000000, None),
            (dt.date(2026, 10, 5), None, "交通費（ETC・ガソリン）", "ETC", "10月前半分", None, 12000),
            (dt.date(2026, 10, 9), "2609-03", "外注費", "サンプル板金", "9月分 支払い", None, 330000),
            (dt.date(2026, 10, 12), None, "消耗品", "ホームセンター", "養生材・工具", None, 8600)]
    for i, r in enumerate(range(C0, C1 + 1)):
        put(ws, f"A{r}", None, "in", DATE)
        put(ws, f"B{r}", None, "in")
        put(ws, f"C{r}", "=" + lookup(f"B{r}", "B"), "auto")
        for col in "DEF":
            put(ws, f"{col}{r}", None, "in")
        put(ws, f"G{r}", None, "in", YEN)
        put(ws, f"H{r}", None, "in", YEN)
        f = f'=IF(COUNT(G{r}:H{r})=0,"",$B$2+SUM(G${C0}:G{r})-SUM(H${C0}:H{r}))'
        put(ws, f"I{r}", f, "auto", YEN)
        if sample and i < len(cash):
            d, a, k, p, m, inn, out = cash[i]
            for col, val in zip("ABDEFGH", [d, a, k, p, m, inn, out]):
                ws[f"{col}{r}"].value = val
    dv_list(ws, f"B{C0}:B{C1}", LED_NO)
    dv_list(ws, f"D{C0}:D{C1}", LST_ACCT)
    ws.freeze_panes = "A6"
    ws.auto_filter.ref = f"A5:I{C1}"
    page(ws)

    # ===================== ダッシュボード =====================
    ws = ws_dash
    widths(ws, {"A": 3, "B": 30, "C": 21, "D": 21, "E": 21, "F": 21, "G": 21, "H": 21})
    ws["B1"].value = "ダッシュボード（全体の状況）"
    ws["B1"].font = Font(name=FONT, size=16, bold=True, color=NAVY)
    ws["B2"].value = "すべて自動計算です。入力は他のシートで行ってください。"
    ws["B2"].font = Font(name=FONT, size=10, color="595959")
    kp = [("現在の残高", f"={q(S_CASH)}!I3", YEN),
          ("未入金の合計（税込）", f"={q(S_LED)}!M3", YEN),
          ("未払いの合計（税込）", f'=SUMIF({q(S_PAY)}!$I${P0}:$I${P1},"未払い",{q(S_PAY)}!$G${P0}:$G${P1})+SUMIF({q(S_PAY)}!$I${P0}:$I${P1},"問題あり",{q(S_PAY)}!$G${P0}:$G${P1})', YEN),
          ("今月が期限の未払い", f'=SUMIFS({q(S_PAY)}!$G${P0}:$G${P1},{q(S_PAY)}!$H${P0}:$H${P1},">="&DATE(YEAR(TODAY()),MONTH(TODAY()),1),{q(S_PAY)}!$H${P0}:$H${P1},"<"&EDATE(DATE(YEAR(TODAY()),MONTH(TODAY()),1),1),{q(S_PAY)}!$I${P0}:$I${P1},"<>支払い済み")', YEN),
          ("期限超過の支払い", f'=COUNTIFS({q(S_PAY)}!$H${P0}:$H${P1},"<"&TODAY(),{q(S_PAY)}!$I${P0}:$I${P1},"<>支払い済み",{q(S_PAY)}!$B${P0}:$B${P1},"<>")', '0"件"'),
          ("粗利の合計（税抜）", f"={q(S_LED)}!J3", YEN)]
    for i, (lab, f, fmt) in enumerate(kp):
        col = "BCDEFG"[i]
        put(ws, f"{col}4", lab, fill=HDR_FILL, bold=True, color="FFFFFF", h="center", wrap=True)
        put(ws, f"{col}5", f, "auto", fmt, bold=True, size=12, h="center")
    ws.row_dimensions[4].height = 32
    ws.row_dimensions[5].height = 36
    ws.conditional_formatting.add("F5", FormulaRule(formula=["F5>0"], font=Font(color="C00000", bold=True), fill=PatternFill("solid", bgColor="F8CBAD")))
    # 状況別
    ws["B7"].value = "■ 案件の状況別"
    ws["B7"].font = Font(name=FONT, bold=True, color=NAVY)
    header(ws, 8, ["状況", "件数", "請負金額（税抜）"], col0=2)
    for i in range(6):
        r = 9 + i
        put(ws, f"B{r}", f"={q(S_SET)}!B{19 + i}", "auto")
        put(ws, f"C{r}", f'=IF(B{r}="","",COUNTIF({led_col("E")},B{r}))', "auto", '0"件"', h="center")
        put(ws, f"D{r}", f'=IF(B{r}="","",SUMIF({led_col("E")},B{r},{led_col("F")}))', "auto", YEN)
    # 月別
    ws["B16"].value = "■ 月別の入出金（開始月は「設定」で変更）"
    ws["B16"].font = Font(name=FONT, bold=True, color=NAVY)
    header(ws, 17, ["月", "入金", "出金", "差引"], col0=2)
    cA = f"{q(S_CASH)}!$A${C0}:$A${C1}"
    for i in range(12):
        r = 18 + i
        put(ws, f"B{r}", f"=EDATE({START},{i})", "auto", 'yyyy"年"m"月"', h="center")
        for col, src in (("C", "G"), ("D", "H")):
            put(ws, f"{col}{r}", f'=SUMIFS({q(S_CASH)}!${src}${C0}:${src}${C1},{cA},">="&B{r},{cA},"<"&EDATE(B{r},1))', "auto", YEN)
        put(ws, f"E{r}", f"=C{r}-D{r}", "auto", YEN)
    put(ws, "B30", "合計", fill=SUB_FILL, bold=True, h="center")
    for col in "CDE":
        put(ws, f"{col}30", f"=SUM({col}18:{col}29)", "auto", YEN, bold=True)
    ch = BarChart()
    ch.type = "col"
    ch.title = "月別の入金・出金"
    ch.height, ch.width = 8.5, 17
    ch.add_data(Reference(ws, min_col=3, max_col=4, min_row=17, max_row=29), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=2, min_row=18, max_row=29))
    ch.y_axis.numFmt = '#,##0'
    ch.y_axis.delete = False
    ch.x_axis.delete = False
    ch.x_axis.number_format = 'm"月"'
    ws.add_chart(ch, "F7")
    # 案件別
    ws["B33"].value = "■ 案件別の収支（台帳の上から15件）"
    ws["B33"].font = Font(name=FONT, bold=True, color=NAVY)
    header(ws, 34, ["案件名", "状況", "請負金額", "外注費", "粗利", "粗利率", "未入金額"], col0=2)
    for i in range(15):
        r = 35 + i
        t = T0 + i
        L = q(S_LED)
        put(ws, f"B{r}", f'=IF({L}!B{t}="","",{L}!B{t})', "auto")
        put(ws, f"C{r}", f'=IF({L}!B{t}="","",{L}!E{t}&"")', "auto", h="center")
        put(ws, f"D{r}", f'=IF({L}!B{t}="","",{L}!F{t})', "auto", YEN)
        put(ws, f"E{r}", f'=IF({L}!B{t}="","",{L}!I{t})', "auto", YEN)
        put(ws, f"F{r}", f'=IF({L}!B{t}="","",{L}!J{t})', "auto", YEN)
        put(ws, f"G{r}", f'=IF({L}!B{t}="","",{L}!K{t})', "auto", "0.0%")
        put(ws, f"H{r}", f'=IF({L}!B{t}="","",{L}!M{t})', "auto", YEN)
    page(ws)

    # ===================== 使い方 =====================
    ws = ws0
    widths(ws, {"A": 3, "B": 22, "C": 90})
    ws["B1"].value = "事務管理ブック" + ("（サンプル入り版）" if sample else "")
    ws["B1"].font = Font(name=FONT, size=18, bold=True, color=NAVY)
    ws["B2"].value = ("※ このファイルの会社名・金額・取引先はすべてサンプルです。実際の運用前に、サンプル行を消してから使ってください。"
                      if sample else "予算・支払い・入出金・見積・請求書を1つにまとめた管理ブックです。")
    ws["B2"].font = Font(name=FONT, size=10, color="C00000" if sample else "595959")
    put(ws, "B4", "色の約束", fill=HDR_FILL, bold=True, color="FFFFFF")
    put(ws, "C4", None, fill=HDR_FILL)
    put(ws, "B5", "黄色のセル", "in", bold=True)
    put(ws, "C5", "自分で入力する場所です。")
    put(ws, "B6", "灰色のセル", "auto", bold=True)
    put(ws, "C6", "自動で計算される場所です。上書きしないでください（数式が消えます）。")
    put(ws, "B8", "最初にやること", fill=HDR_FILL, bold=True, color="FFFFFF")
    put(ws, "C8", None, fill=HDR_FILL)
    steps = [
        ("① 設定", "「設定」シートに自社名・住所・インボイス番号・振込先を入れます（見積書・請求書に自動で入ります）。"),
        ("② 案件台帳", "案件を1行ずつ登録します（案件No・案件名・取引先・請負金額）。案件Noは自由な番号でOK（例: 2610-01）。"),
        ("③ 見積書", "案件Noを選ぶと宛名・件名が入ります。明細（工種・品名・数量・単価）を入力。印刷はA3、1ページに収まります。"),
        ("④ 予算管理", "業者に発注したら1件1行で入力。案件ごとの外注費と粗利が「案件台帳」に自動で出ます。"),
        ("⑤ 請求書", "案件Noを選ぶと、請負金額が1行（一式）で入ります。分割請求は「請求割合」を変えます（例: 50%）。"),
        ("⑥ 支払管理", "業者から届いた請求書を入力。支払状況を選ぶと色が変わり、期限切れは赤字になります。"),
        ("⑦ 入出金帳", "入金・出金を1行ずつ入力。残高が自動で出ます。案件Noを選ぶと、その案件の入金済額・未入金額が台帳に出ます。"),
        ("⑧ ダッシュボード", "残高・未入金・未払い・月別の入出金・案件別の粗利を一覧で見られます（入力は不要）。"),
    ]
    for i, (a, b) in enumerate(steps):
        r = 9 + i
        put(ws, f"B{r}", a, bold=True, fill=SUB_FILL)
        put(ws, f"C{r}", b, wrap=True)
        ws.row_dimensions[r].height = 32
    put(ws, "B18", "気をつけること", fill=HDR_FILL, bold=True, color="FFFFFF")
    put(ws, "C18", None, fill=HDR_FILL)
    notes = [
        "請負金額・見積金額は、必ず本人が最終確認して入力してください（自動では決まりません）。",
        "行を並べ替える時は、同じシートの行ごと選んで「データ」タブの並べ替えを使ってください。",
        "案件台帳の合計・粗利は、状況が「見積中」「失注」の案件を除いて計算します（受注した案件だけを数えるため）。",
        "入出金帳の残高は上から順に計算します。行の途中に空き行を作らないでください。",
        "行が足りなくなったら、最終行の1行上を選んで行を挿入し、数式のある行をコピーしてください。",
        "他の人に渡す時は「設定」の自社情報を書き換え、サンプルを消してお渡しください（空白版もあります）。",
        "マクロは使っていません。ExcelでもGoogleスプレッドシートでも開けます。",
    ]
    for i, t in enumerate(notes):
        r = 19 + i
        put(ws, f"B{r}", "・", h="right")
        put(ws, f"C{r}", t, wrap=True)
        ws.row_dimensions[r].height = 28
    page(ws, landscape=False)

    wb.active = 0
    wb.save(path)


if __name__ == "__main__":
    out = sys.argv[1]
    build(sample=(sys.argv[2] == "sample"), path=out)
    print("saved", out)
